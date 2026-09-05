"""Reads/writes the jobs workbook -- the single source of truth for
"have I already logged this job" and where "Is Applied" gets toggled.

Column order matches exactly what was asked for: Company, Job Title,
Salary, Requirements, Apply Link, Is Applied -- with Source and
Location appended after as useful extra context, not swapped in
anywhere that would change the requested layout.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Job

SHEET_NAME = "Jobs"
HEADERS = ["Company", "Job Title", "Salary", "Requirements", "Apply Link", "Is Applied", "Source", "Location"]
_MAX_ROWS_FOR_VALIDATION = 100_000


def _new_workbook() -> WorkbookType:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    ws.freeze_panes = "A2"
    for column_cells in ws.columns:
        ws.column_dimensions[column_cells[0].column_letter].width = 22

    # "Is Applied" as a real Yes/No dropdown instead of free text --
    # column F is the 6th header, matching HEADERS above.
    validation = DataValidation(type="list", formula1='"No,Yes"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"F2:F{_MAX_ROWS_FOR_VALIDATION}")
    return wb


def load_or_create(path: Path) -> WorkbookType:
    if path.exists():
        return load_workbook(path)
    return _new_workbook()


def _worksheet(wb: WorkbookType):
    return wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active


def existing_keys(wb: WorkbookType, dedup_strategy: str) -> set[str]:
    """Every dedup key already present in the sheet, built from a
    throwaway Job for each row so the key logic lives in exactly one
    place (Job.dedup_key), not duplicated here."""
    ws = _worksheet(wb)
    keys = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not any(row):
            continue
        company, title, _salary, _requirements, link = (row + (None,) * 5)[:5]
        if not company and not title:
            continue
        placeholder = Job(source="", company=company or "", title=title or "", apply_link=link or "")
        keys.add(placeholder.dedup_key(dedup_strategy))
    return keys


def append_jobs(wb: WorkbookType, jobs: list[Job]) -> int:
    ws = _worksheet(wb)
    for job in jobs:
        ws.append([
            job.company,
            job.title,
            job.salary_raw,
            job.requirements,
            job.apply_link,
            "No",
            job.source,
            job.location,
        ])
    return len(jobs)


def save(wb: WorkbookType, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
